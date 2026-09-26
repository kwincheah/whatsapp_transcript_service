import io

import pytest
from docx import Document

from app import pipeline
from app.documents import DocumentError, extract_text, kind_of
from tests.conftest import FakeAI


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid PDF with one line of Helvetica text per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_id = len(objs)
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>"
        )
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{o}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def make_docx() -> bytes:
    d = Document()
    d.add_paragraph("Contract between A and B")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Fee", "RM 5,000"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_kind_detection():
    assert kind_of("application/pdf", "x") == "PDF"
    assert kind_of("", "Report.PDF") == "PDF"
    assert kind_of("application/octet-stream", "notes.md") == "Text"
    assert kind_of("application/zip", "a.zip") is None


def test_extract_pdf_docx_text():
    ex = extract_text(make_pdf(["Hello page one", "Second page"]), "application/pdf", "a.pdf", 10_000)
    assert ex.kind == "PDF" and ex.pages == 2
    assert "Hello page one" in ex.text and "Second page" in ex.text
    ex = extract_text(make_docx(), "", "c.docx", 10_000)
    assert "Contract between A and B" in ex.text and "Fee | RM 5,000" in ex.text
    assert extract_text("héllo".encode(), "text/plain", "a.txt", 100).text == "héllo"
    with pytest.raises(DocumentError):
        extract_text(b"garbage", "application/pdf", "bad.pdf", 100)
    with pytest.raises(DocumentError):
        extract_text(b"x", "application/zip", "a.zip", 100)


async def test_document_reply_with_question(ctx_factory):
    ctx, sent = ctx_factory()
    await pipeline.handle_document(ctx, make_pdf(["Revenue grew 10%"]), "application/pdf", "q3.pdf", "Did revenue grow?")
    call = ctx.ai.doc_calls[0]
    assert call["question"] == "Did revenue grow?" and call["truncated"] is False
    out = sent[-1]
    assert out.startswith("📄 *Quarterly Report*\nq3.pdf · PDF · 1 page")
    assert "❓ *Did revenue grow?*\nYes, 10%." in out
    assert "💡 *Summary*\nRevenue grew." in out and "• Revenue +10%" in out
    assert "Analysis deepseek-v4-flash 1.20s" in out and "12.0k tokens in" in out and "$0.0040" in out
    assert ctx.store.search("60111", "revenue")[0].kind == "document"


async def test_long_document_truncated_with_ack(ctx_factory):
    ctx, sent = ctx_factory(doc_max_input_chars=30_000)
    text = ("lorem ipsum dolor " * 5000).encode()
    await pipeline.handle_document(ctx, text, "text/plain", "big.txt", None)
    assert ctx.ai.doc_calls[0]["truncated"] is True
    assert len(ctx.ai.doc_calls[0]["text"]) == 30_000
    assert sent[0].startswith("📄 Reading *big.txt*")
    assert "⚠️ Long document: analysed the first ~33%." in sent[-1]


async def test_empty_and_unsupported_documents(ctx_factory):
    ctx, sent = ctx_factory()
    await pipeline.handle_document(ctx, make_pdf([""]), "application/pdf", "scan.pdf", None)
    assert "couldn't find any text" in sent[-1] and ctx.ai.doc_calls == []
    await pipeline.handle_document(ctx, b"PK..", "application/zip", "a.zip", None)
    assert "I can read PDF" in sent[-1]
